"""Запись «Факт»-сумм в шаблон финплана (openpyxl).

Особенности шаблона (выяснены по файлу пользователя):
  - месяцы — объединённые пары колонок «План|Факт», строка 2 = имя месяца,
    строка 3 = План/Факт; «Факт» = стартовая колонка + 1;
  - листовые «Факт»-ячеек уже содержат ручные накопительные формулы
    вида =20800+4200-... — мы ДОПИСЫВАЕМ '+сумма', а не затираем ячейку;
  - расходы в шаблоне хранятся положительными числами;
  - исходный шаблон не перезаписывается: сохраняем новую копию в output/.
"""
import datetime as dt
import os

import openpyxl
from openpyxl.utils import get_column_letter

HEADER_ROW_MONTH, HEADER_ROW_PF = 2, 3


def find_month_columns(ws) -> dict[str, int]:
    """{'Сентябрь 2026': колонка_План} по строке 2."""
    out = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=HEADER_ROW_MONTH, column=col).value
        if v and "Итого" not in str(v) and "2026" in str(v):
            out[str(v)] = col
    return out


def write_facts(template_path: str, out_dir: str, month_key: str,
                by_category: dict[str, float], cats: list[dict]) -> tuple[str, list[dict]]:
    """Дописывает суммы в колонку «Факт» месяца. Возвращает (путь, журнал записей)."""
    wb = openpyxl.load_workbook(template_path)  # формулы сохраняются
    ws = wb["Финансовый план"]
    months = find_month_columns(ws)
    if month_key not in months:
        raise KeyError(f"В шаблоне нет месяца «{month_key}». Есть: {list(months)}")
    fact_col = months[month_key] + 1

    row_by_name = {c["name"]: (c["row"], c["type"]) for c in cats}
    journal = []
    for name, value in by_category.items():
        row, kind = row_by_name[name]
        write_value = abs(value) if kind == "expense" else value
        write_value = round(write_value, 2)
        if abs(write_value) < 0.005:
            continue
        cell = ws.cell(row=row, column=fact_col)
        old = cell.value
        if old is None:
            cell.value = write_value
        elif isinstance(old, str) and old.startswith("="):
            cell.value = f"{old}+{write_value}"      # дописываем к формуле-накопителю
        elif isinstance(old, (int, float)):
            cell.value = f"={old}+{write_value}"     # константу превращаем в формулу
        else:
            raise ValueError(f"Неожиданное содержимое {cell.coordinate}: {old!r}")
        journal.append({"row": row, "category": name,
                        "col": f"{get_column_letter(fact_col)}{row}",
                        "old": old, "added": write_value})

    os.makedirs(out_dir, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = os.path.join(out_dir, f"финплан_обновлён_{month_key.replace(' ', '_')}_{stamp}.xlsx")
    wb.save(out_path)
    return out_path, journal


def read_plan(template_path: str, month_key: str, cats: list[dict]) -> dict[str, float]:
    """Плановые суммы месяца (для сравнения план/факт в комментарии).
    Читает кэшированные значения (data_only), формулы не вычисляет."""
    wb = openpyxl.load_workbook(template_path, data_only=True)
    ws = wb["Финансовый план"]
    months = find_month_columns(ws)
    if month_key not in months:
        return {}
    plan_col = months[month_key]
    return {c["name"]: ws.cell(row=c["row"], column=plan_col).value
            for c in cats
            if isinstance(ws.cell(row=c["row"], column=plan_col).value, (int, float))}


def load_balance_rows(path: str | None = None) -> list[dict]:
    """Строки остатков из data/balance_rows.json (секции 1 и 5 шаблона)."""
    import json as _json
    if path is None:
        path = os.path.join(os.path.dirname(__file__), "..", "data",
                            "balance_rows.json")
    with open(path, encoding="utf-8") as f:
        return _json.load(f)["balance_rows"]


def write_balances(target_path: str, out_dir: str, month_key: str,
                   entries: list[dict]) -> tuple[str, list[dict]]:
    """Записывает остатки в колонку «Факт» месяца.

    ВАЖНО: остаток — снимок состояния, поэтому ячейка ПЕРЕЗАПИСЫВАЕТСЯ
    (в отличие от оборотов write_facts, которые дописываются).
    entries: [{'row': int, 'value': float}]
    """
    wb = openpyxl.load_workbook(target_path)
    ws = wb["Финансовый план"]
    months = find_month_columns(ws)
    if month_key not in months:
        raise KeyError(f"В шаблоне нет месяца «{month_key}». Есть: {list(months)}")
    fact_col = months[month_key] + 1

    journal = []
    for e in entries:
        cell = ws.cell(row=e["row"], column=fact_col)
        old = cell.value
        cell.value = round(float(e["value"]), 2)
        journal.append({"row": e["row"], "col": f"{get_column_letter(fact_col)}{e['row']}",
                        "old": old, "new": cell.value})

    os.makedirs(out_dir, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = os.path.join(out_dir, f"финплан_остатки_{month_key.replace(' ', '_')}_{stamp}.xlsx")
    wb.save(out_path)
    return out_path, journal
