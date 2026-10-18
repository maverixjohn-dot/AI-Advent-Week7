"""Таксономия категорий и маппинг операций.

Источники данных (всё правится пользователем, без кода):
  data/template.xlsx  — сам шаблон финплана (структура ниже);
  data/categories.json — сгенерировано из шаблона (tools/rebuild_categories.py);
  data/members.json    — псевдонимы членов семьи для привязки выписок
                         (поставляется ПУСТЫМ — обезличено);
  лист «Маппинг» в template.xlsx (необязательно): A=ключевое слово,
                         B=статья — подсказки классификатору.

Структура шаблона, от которой зависит генерация:
  строка 2 — названия месяцев (объединённые пары План|Факт);
  колонка A — секции («2. Доходы», «3. Расходы»…), B — подсекции,
  C — листовые строки (статьи); лист — «Финансовый план».
"""
import json
import os
import re

_DATA = os.path.join(os.path.dirname(__file__), "..", "data")


def _path(name, override=None):
    return override or os.path.join(_DATA, name)


def load_categories(path: str | None = None) -> list[dict]:
    """[{'row', 'name', 'type': income|expense}]"""
    with open(_path("categories.json", path), encoding="utf-8") as f:
        return json.load(f)["categories"]


def load_members(path: str | None = None) -> list[str]:
    """Псевдонимы членов семьи. Пустой список = режим без привязки."""
    p = _path("members.json", path)
    if not os.path.exists(p):
        return []
    with open(p, encoding="utf-8") as f:
        return json.load(f).get("members", [])


def member_suffix(name: str, members: list[str]) -> str | None:
    """Если статья заканчивается на псевдоним члена — он и есть владелец."""
    for m in members:
        if name.endswith(m):
            return m
    return None


def allowed_names(cats: list[dict], owner: str | None,
                  members: list[str]) -> list[str]:
    """Категории, доступные для владельца owner:
    нейтральные (без суффикса члена) + его собственные.
    Пустой members или owner=None — весь список (режим без привязки)."""
    if not members or not owner:
        return [c["name"] for c in cats]
    out = []
    for c in cats:
        suffix = member_suffix(c["name"], members)
        if suffix is None or suffix == owner:
            out.append(c["name"])
    return out


def resolve_owner_variant(base_name: str, owner: str | None,
                          allowed: list[str], members: list[str]) -> str | None:
    """Для базового имени («% по вкладам») находит вариант владельца
    («% по вкладам <owner>») среди разрешённых; иначе None."""
    if not owner or not members:
        return None
    candidate = f"{base_name} {owner}"
    return candidate if candidate in allowed else None


def category_names(cats: list[dict], kind: str | None = None) -> list[str]:
    return [c["name"] for c in cats if kind is None or c["type"] == kind]


def fallback_category(cats: list[dict], kind: str) -> str:
    """Категория-приёмник для нераспознанного: 'Прочее' / 'Прочие расходы'."""
    for c in cats:
        if c["type"] == kind and c["name"].lower().startswith("проч"):
            return c["name"]
    raise KeyError(f"Нет категории-приёмника для типа {kind}")


def load_mapping(template_path: str) -> dict[str, str]:
    """Лист «Маппинг» шаблона: A=ключевое слово (в нижнем регистре),
    B=статья (точное имя строки финплана). Пусто, если листа нет."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(template_path, read_only=True)
    except Exception:
        return {}
    if "Маппинг" not in wb.sheetnames:
        return {}
    ws = wb["Маппинг"]
    out = {}
    for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        kw, cat = (row + (None, None))[:2]
        if kw and cat and str(kw).strip().lower() != "ключевое слово":
            out[str(kw).strip().lower()] = str(cat).strip()
    wb.close()
    return out


def apply_mapping(description: str, mapping: dict[str, str],
                  allowed: list[str]) -> str | None:
    """Детерминированное правило «ключевое слово -> статья» (лист «Маппинг»).

    Применяется КОДОМ до вызова LLM: LLM не может проигнорировать правило.
    Совпадение — подстрока без учёта регистра; более длинные ключи
    приоритетнее коротких. Статья должна быть среди разрешённых для
    владельца — иначе None (решит классификатор)."""
    if not mapping:
        return None
    desc = description.lower()
    for kw in sorted(mapping, key=len, reverse=True):
        if kw in desc:
            target = mapping[kw]
            return target if target in allowed else None
    return None


def validate_mapping(mapping: dict[str, str], cats: list[dict]) -> list[str]:
    """Статьи из маппинга, которых НЕТ в таксономии шаблона.
    Такие правила бесполезны: статью надо добавить в шаблон
    и пересобрать categories.json (tools/rebuild_categories.py)."""
    names = {c["name"] for c in cats}
    return sorted({v for v in mapping.values() if v not in names})
