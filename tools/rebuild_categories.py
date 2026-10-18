"""Регенерация categories.json и balance_rows.json из data/template.xlsx.

Запускать после ЛЮБОГО изменения шаблона (добавили/переименовали статьи):
    python3 tools/rebuild_categories.py

Структура, которую скрипт ожидает:
  лист «Финансовый план»;
  строка 2 — названия месяцев; строка 3 — План/Факт;
  колонка A — секции («1.», «2.», «3.», «5.»…), B — подсекции («2.1.»…),
  C — листовые строки (статьи). Доходы = секция 2, расходы = секция 3,
  остатки — листья секций 1 и 5.
"""
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(BASE, "data", "template.xlsx")
SECTION_RE = re.compile(r"^(\d)\.")


def main():
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["Финансовый план"]

    cats, current = [], None
    for r in range(4, 84):
        a, b, c = (ws.cell(row=r, column=i).value for i in (1, 2, 3))
        if a:
            m = SECTION_RE.match(str(a).strip())
            if m:
                current = {"2": "income", "3": "expense"}.get(m.group(1))
                continue
        if b and not c and re.match(r"^[23]\.\d\.", str(b)):
            continue
        if c and current in ("income", "expense"):
            cats.append({"row": r, "name": str(c).strip(), "type": current})

    balance_rows = []
    for r in list(range(6, 15)) + list(range(62, 71)):
        name = ws.cell(row=r, column=3).value
        if name:
            section = "начальное сальдо (1)" if r < 60 else "конечное сальдо (5)"
            balance_rows.append({"row": r, "name": str(name).strip(),
                                 "section": section})

    with open(os.path.join(BASE, "data", "categories.json"), "w",
              encoding="utf-8") as f:
        json.dump({"categories": cats}, f, ensure_ascii=False, indent=2)
    with open(os.path.join(BASE, "data", "balance_rows.json"), "w",
              encoding="utf-8") as f:
        json.dump({"balance_rows": balance_rows}, f, ensure_ascii=False, indent=2)
    print(f"categories.json: {len(cats)} статей "
          f"({sum(1 for c in cats if c['type']=='income')} доходных, "
          f"{sum(1 for c in cats if c['type']=='expense')} расходных)")
    print(f"balance_rows.json: {len(balance_rows)} строк остатков")
    print("members.json не трогаю — он ваш, локальный.")


if __name__ == "__main__":
    main()
